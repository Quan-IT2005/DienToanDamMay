import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document(r'C:\Users\quan\Downloads\Nhóm 02 - Báo Cáo Buổi 03.docx')
print('Total paragraphs:', len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._r.xml for r in p.runs)
    if has_img:
        prev_p = doc.paragraphs[i-1].text.strip() if i > 0 else ''
        print(f'P{i} [IMG] after: "{prev_p[:60]}"')
